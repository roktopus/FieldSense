#!/usr/bin/env python3
"""Build the two-page FieldSense write-up (system design + results)."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn, nsmap
from docx.shared import Inches, Pt, RGBColor, Twips, Cm

ROOT = Path(__file__).resolve().parents[1]
ARCH = ROOT / "docs" / "assets" / "architecture.png"
OUT_DOCX = ROOT / "docs" / "FieldSense_Writeup_Design_and_Results.docx"
OUT_DOC = ROOT / "docs" / "FieldSense_Writeup_Design_and_Results.doc"


def set_run(run, *, size=9, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_para(doc, text, *, size=9, bold=False, italic=False, space_after=4, space_before=0, color=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    run = p.add_run(text)
    set_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_heading_custom(doc, text, size=12):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(6)
    pf.space_after = Pt(4)
    run = p.add_run(text)
    set_run(run, size=size, bold=True, color=RGBColor(0x0B, 0x1F, 0x33))
    return p


def add_bullet(doc, text, *, size=9):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.space_after = Pt(2)
    pf.space_before = Pt(0)
    pf.line_spacing = 1.05
    # clear default run and set ours
    if p.runs:
        p.runs[0].text = text
        set_run(p.runs[0], size=size)
    else:
        run = p.add_run(text)
        set_run(run, size=size)
    return p


def set_narrow_margins(section):
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.6)
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)


def add_page_number(section):
    footer = section.footer
    footer.is_linked_to_previous = False
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("FieldSense · FlytBase Hacktoberfest 2026 · Page ")
    set_run(run, size=8, color=RGBColor(0x64, 0x74, 0x8B))

    # PAGE field
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r2 = p.add_run()
    r2._r.append(fld_begin)
    r2._r.append(instr)
    r2._r.append(fld_end)
    set_run(r2, size=8, color=RGBColor(0x64, 0x74, 0x8B))

    run3 = p.add_run(" of 2")
    set_run(run3, size=8, color=RGBColor(0x64, 0x74, 0x8B))


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    set_narrow_margins(section)
    add_page_number(section)

    # -------- PAGE 1: System design --------
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    r = title.add_run("FieldSense — System Design Write-up")
    set_run(r, size=16, bold=True, color=RGBColor(0x0B, 0x1F, 0x33))

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.paragraph_format.space_after = Pt(6)
    r = sub.add_run(
        "Asset Intelligence Copilot · Page 1 of 2 · Architecture & operating model"
    )
    set_run(r, size=9, italic=True, color=RGBColor(0x47, 0x55, 0x69))

    if ARCH.exists():
        doc.add_picture(str(ARCH), width=Inches(6.8))
        last = doc.paragraphs[-1]
        last.alignment = WD_ALIGN_PARAGRAPH.CENTER
        last.paragraph_format.space_after = Pt(4)
        cap = add_para(
            doc,
            "Figure 1. FieldSense modular monolith — store → retrieve → tools → answer envelope.",
            size=8,
            italic=True,
            space_after=6,
            color=RGBColor(0x64, 0x74, 0x8B),
        )
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    add_heading_custom(doc, "1. How data is stored, indexed, and retrieved", size=11)
    add_para(
        doc,
        "FieldSense is a modular monolith: React/Vite UI, FastAPI API, background worker, and "
        "PostgreSQL 16 with pgvector. Challenge sample files mount read-only; an idempotent "
        "ingestion CLI loads wells, areas, weekly inspections, maintenance/repairs, regulatory "
        "events, source notes, manuals/SOPs, and extracted report facts into normalized tables. "
        "Content hashes keep re-imports safe. Well counts and area names are discovered at import — "
        "not hardcoded.",
        size=9,
        space_after=3,
    )
    add_bullet(
        doc,
        "Store: relational tables for assets/events; corrections, assertions, and rule overrides "
        "as append-only knowledge; result sets + evidence items per answer.",
    )
    add_bullet(
        doc,
        "Index / retrieve: primary path is typed SQL and deterministic domain services keyed by "
        "asset_id, area, and as-of time. Lexical full-text search is the always-on document "
        "fallback; embeddings/pgvector are optional. Evaluation answer sheets are never indexed.",
    )
    add_bullet(
        doc,
        "Clock: analysis uses a controllable as-of (default 2026-10-17T00:00:00Z). "
        "“Right now” = latest inspection at/before that clock.",
    )

    add_heading_custom(doc, "2. Tools the copilot uses", size=11)
    add_para(
        doc,
        "A single orchestrator routes each question (deterministic intent router; optional "
        "OpenAI-compatible LLM for planning/explanation when configured). Typed tools include:",
        size=9,
        space_after=3,
    )
    add_bullet(
        doc,
        "Analytics: well status/ΔT, area non-producing tables, temperature series, downtime "
        "rankings, spill history/map/heatmap, repair costs.",
    )
    add_bullet(
        doc,
        "Compliance: SOP breach assessment (SR-01/02/03, RC-04, RC-02), containment-hour facts "
        "from repair reports, pump-jack RCA thresholds from manuals + team overrides.",
    )
    add_bullet(
        doc,
        "Feedback: measurement/detection corrections (preview → apply), shared notes, "
        "unconfirmed claims, well-type rule overrides.",
    )
    add_bullet(
        doc,
        "UI schema tools: status card, table, bar/line chart, map, heatmap — bound to server "
        "result rows so visuals cannot invent numbers.",
    )

    add_heading_custom(doc, "3. Combining numbers, documents, and notes", size=11)
    add_para(
        doc,
        "Numbers always come from deterministic code over an effective-inspection layer: "
        "production status from wellhead−air ΔT vs T200 (15 °C) / T350 (20 °C) manual thresholds; "
        "downtime = 7 × non-producing weekly inspections in the window; spill volumes and "
        "exclusions from detections ± formal corrections/notes. Manuals and SOPs seed domain "
        "rules (reporting windows, containment limits, RCA counts). Free-text repair reports "
        "contribute extracted facts (e.g. containment hours) with coverage when missing. "
        "Human notes explain context and conflicts but never silently override measurements. "
        "Every factual answer carries citations and resolvable evidence IDs "
        "(INS-…, MNT-…, CTX-…, REG-…).",
        size=9,
        space_after=3,
    )

    add_heading_custom(doc, "4. Two memory levels", size=11)
    add_bullet(
        doc,
        "Private memory: conversations and ConversationReference (discussed well sets) are "
        "scoped by demo user id in the backend before retrieval. “Wells we discussed yesterday” "
        "only sees that user’s multi-well history — never another operator’s chat.",
    )
    add_bullet(
        doc,
        "Shared knowledge: formal corrections, FYI/planned-work assertions, and well-type rule "
        "overrides are visible to everyone and immediately affect analytics via the effective layer.",
    )

    add_heading_custom(doc, "5. Conflicts and new data", size=11)
    add_para(
        doc,
        "Conflicts are disclosed, not resolved quietly: inspector notes vs corrected readings, "
        "unconfirmed “running fine / back online” claims vs latest ΔT, and informal disagreeing "
        "notes on a spill (Q18) keep the detection if no formal correction exists. Corrections "
        "preserve originals, record author/reason, and supersede prior actives on the same field. "
        "New dataset versions import idempotently alongside existing chats/corrections; "
        "quarantine paths avoid silent row drops.",
        size=9,
        space_after=3,
    )

    add_heading_custom(doc, "6. Models used and cost", size=11)
    add_para(
        doc,
        "Hack Day path: no paid LLM required. The deterministic router answers Q1–Q18 analytical "
        "and memory scenarios offline (≈ $0 inference). Optional OpenAI-compatible endpoint "
        "(OpenRouter / Ollama / Google AI Studio) can be enabled via MODEL_* env vars for richer "
        "phrasing; tool-calling models are preferred. Embeddings remain optional. Hosting cost "
        "for the demo is a single Docker Compose stack (Postgres + API + worker + nginx).",
        size=9,
        space_after=3,
    )

    add_heading_custom(doc, "7. What we would do next", size=11)
    add_bullet(doc, "Replace brittle phrase routing with a small tool-calling planner + JSON schemas.")
    add_bullet(doc, "Harden report-fact extraction (LLM-assisted, human-confirm) for SOP narratives.")
    add_bullet(doc, "SSE streaming polish, Playwright e2e for Q1–Q18, and production auth beyond demo X-User-Id.")
    add_bullet(doc, "Richer conflict UX (side-by-side evidence) and continuous re-eval on the final test set.")

    # -------- PAGE 2: Results --------
    doc.add_page_break()

    title2 = doc.add_paragraph()
    title2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title2.paragraph_format.space_after = Pt(2)
    r = title2.add_run("FieldSense — Results Write-up")
    set_run(r, size=16, bold=True, color=RGBColor(0x0B, 0x1F, 0x33))

    sub2 = doc.add_paragraph()
    sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub2.paragraph_format.space_after = Pt(8)
    r = sub2.add_run(
        "Page 2 of 2 · Example scenarios Q1–Q18 on sample-v1 (as-of 17 Oct 2026) · "
        "Screenshots in FieldSense_Q1-Q18_System_Answers.doc"
    )
    set_run(r, size=9, italic=True, color=RGBColor(0x47, 0x55, 0x69))

    add_heading_custom(doc, "1. Scope of evaluation", size=11)
    add_para(
        doc,
        "We ran the official example scenarios end-to-end in the Docker UI with the sample "
        "dataset (200 wells). “Today” / as-of = 17 Oct 2026. Part A used Michael Turner; "
        "Part B switched Jessica Reed, Laura Simmons, Jake Morales, and Tyler Hughes. "
        "Full answer text, tables, and UI screenshots are in the companion document "
        "FieldSense_Q1-Q18_System_Answers.doc. Below: what matched the expected sheets, "
        "what was weaker, and why.",
        size=9,
        space_after=4,
    )

    add_heading_custom(doc, "2. What worked", size=11)
    add_bullet(
        doc,
        "Q1–Q5 (single-well analytics): Correct not-producing status for WELL-0001 "
        "(INS-0001-0203, Ryan Brooks, ΔT 3.2 °C, stale “active”); 13-point temperature series "
        "with 8 Oct / 15 Oct offline marks; 6 spills / 15.1 bbl with rainwater INS-0001-0070 "
        "excluded; 21 downtime days in 2025 with maintenance causes.",
    )
    add_bullet(
        doc,
        "Q2 / Q6–Q8 (area visuals): Four non-producing NE Block 8 wells in a table; top "
        "downtime bar chart WELL-0121 (49), WELL-0107/0138 (42); four real-spill map pins; "
        "2026 area×month heatmap with formal false-positive corrections removed.",
    )
    add_bullet(
        doc,
        "Q10 containment: WELL-0082 and WELL-0100 flagged over the 4-hour SOP limit from "
        "report-extracted hours.",
    )
    add_bullet(
        doc,
        "Q11 private memory: After Q6, repair costs returned WELL-0121 USD 25,794 · "
        "WELL-0138 USD 46,942 · WELL-0107 USD 5,176 from Michael’s discussion only.",
    )
    add_bullet(
        doc,
        "Q12–Q15 corrections (shared): Thermal 41.0 °C on INS-0164-0205 → producing, conflict "
        "with “Pump jack not moving” disclosed; Q2 shrinks to three wells; spill 5.4 bbl on "
        "INS-0151-0201 → reportable SR-01/SR-03 narrative; Jessica and Laura see the same "
        "shared corrections.",
    )
    add_bullet(
        doc,
        "Q14 team rule: Baseline RCA trio; after T350-after-2 override, WELL-0120 appears and "
        "T200 WELL-0102 does not; provenance cites the team rule.",
    )
    add_bullet(
        doc,
        "Q16–Q18 multi-user: Jake’s FYI shared to Michael; Jessica’s “three wells yesterday” "
        "correctly refuses Michael’s history; unconfirmed claims vs inspections recorded; "
        "Q18 keeps the oil detection while showing disagreeing notes.",
    )

    add_heading_custom(doc, "3. What didn’t (or was weaker) — and why", size=11)
    add_bullet(
        doc,
        "Q9 compliance narrative density: We surface SR-01/02/03, RC-04, USD 5,000 fines, "
        "pending RCA/records, and met RC-02, but the breach list can over-count related "
        "regulatory rows versus the brief’s tidy “four breaches” storyline. Cause: event-level "
        "enumeration vs curated spill-centric narrative. Fix: group breaches by initiating spill "
        "and separate “still open” fines more clearly in the summary.",
    )
    add_bullet(
        doc,
        "Intent routing brittleness: Early builds mis-routed “not producing oil right now” "
        "(substring of single-well status) and “heatmap” (substring “map”). Fixed with ordered "
        "intents, but paraphrase coverage still depends on keywords rather than a planner. "
        "Why it matters: final Hack Day questions may not match example phrasing.",
    )
    add_bullet(
        doc,
        "Report-fact / SOP depth: Containment hours work when regex/heuristics extract them; "
        "richer SR-02 storytelling from full report prose is still thinner than a human reader. "
        "Maps need network tiles (table fallback offline).",
    )
    add_bullet(
        doc,
        "No LLM cost, but also less free-form prose: Answers are accurate and evidenced; "
        "they are less essay-like than a large model might write. Optional MODEL_* wiring is "
        "ready when a provider is available.",
    )
    add_bullet(
        doc,
        "Final sealed test set: Released only in the last window of the hack day — not "
        "pre-loaded here. We validated against the published example scenarios and kept eval "
        "fixtures out of the retrieval index so the same stack can run the finals cleanly.",
    )

    add_heading_custom(doc, "4. Evidence that the visuals matched the ask", size=11)
    add_para(
        doc,
        "Status card (Q1/Q15), tables (Q2/Q4/Q9/Q10/Q11/Q14), line chart with offline marks (Q3), "
        "bar chart (Q6), map (Q7), heatmap (Q8), correction receipts (Q12/Q13), conflict reviews "
        "(Q17/Q18). Where the question named a visual, the UI rendered that kind; where it did "
        "not (Q5), we returned a short evidenced answer with no chart.",
        size=9,
        space_after=4,
    )

    add_heading_custom(doc, "5. Bottom line", size=11)
    add_para(
        doc,
        "FieldSense meets the challenge’s core bar: deterministic, evidence-backed answers over "
        "real sample data; generative UI bound to server rows; shared corrections with private "
        "chat isolation; honest conflict handling. Remaining gaps are narrative polish on "
        "compliance, planner generality for unseen phrasings, and deeper report understanding — "
        "not missing foundations. Companion artifacts: architecture diagram in README/docs, "
        "Q1–Q18 screenshot pack, and ADRs under docs/decisions/.",
        size=9,
        space_after=2,
    )

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOCX))
    doc.save(str(OUT_DOC))
    print(f"Wrote {OUT_DOCX}")
    print(f"Wrote {OUT_DOC}")


if __name__ == "__main__":
    build()
