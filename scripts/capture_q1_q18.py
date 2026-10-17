#!/usr/bin/env python3
"""Capture Q1–Q18 from the live FieldSense UI and build a Word report."""

from __future__ import annotations

import json
import re
from pathlib import Path

import httpx
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from playwright.sync_api import sync_playwright

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "scenario-capture"
SHOTS = OUT / "screenshots"
API = "http://localhost:8000"
UI = "http://localhost:5173"

USERS = {
    "michael": ("U-201", "Michael Turner"),
    "jessica": ("U-202", "Jessica Reed"),
    "jake": ("U-102", "Jake Morales"),
    "tyler": ("U-204", "Tyler Hughes"),
    "laura": ("U-401", "Laura Simmons"),
}

SCENARIOS: list[dict] = [
    {
        "id": "Q1",
        "user": "michael",
        "new_chat": True,
        "question": (
            "Is WELL-0001 producing oil right now? When was it last inspected, and by whom? "
            "Show me its status card."
        ),
        "expected_note": (
            "Expected: Not producing. INS-0001-0203 on 15 Oct 2026 by Ryan Brooks (drone). "
            "ΔT 3.2 °C below T200. CTX-0001-003 planned shut-in. Recorded status active (stale)."
        ),
    },
    {
        "id": "Q2",
        "user": "michael",
        "new_chat": True,
        "question": (
            "Which wells in Northeast Block 8 are not producing oil right now, and why? "
            "Show them in a table."
        ),
        "expected_note": (
            "Expected table: WELL-0164, WELL-0170, WELL-0172, WELL-0175 (four T200 wells)."
        ),
    },
    {
        "id": "Q3",
        "user": "michael",
        "new_chat": True,
        "question": (
            "Show a line chart of WELL-0001's wellhead temperature against air temperature "
            "for the last 13 weeks, and mark when it was not producing oil."
        ),
        "expected_note": "Expected: 13 readings; not producing on 8 Oct and 15 Oct 2026.",
    },
    {
        "id": "Q4",
        "user": "michael",
        "new_chat": True,
        "question": (
            "How many oil spills has WELL-0001 had in the last 4 years, and how much oil in total? "
            "List them in a table."
        ),
        "expected_note": "Expected: 6 spills, 15.1 bbl; INS-0001-0070 rainwater excluded.",
    },
    {
        "id": "Q5",
        "user": "michael",
        "new_chat": True,
        "question": "How many days was WELL-0001 down in 2025?",
        "expected_note": "Expected: 21 days from 3 non-producing inspections; no visual required.",
    },
    {
        "id": "Q6",
        "user": "michael",
        "new_chat": True,
        "question": (
            "Which 3 wells in Northeast Block 7 had the most downtime in the last 4 months? "
            "Show a bar chart."
        ),
        "expected_note": "Expected: WELL-0121 49d, WELL-0138/0107 42d (16 Jun–16 Oct 2026).",
    },
    {
        "id": "Q7",
        "user": "michael",
        "new_chat": True,
        "question": (
            "Show on a map every well in Northeast Block 8 with a real spill in the last 90 days."
        ),
        "expected_note": "Expected map: WELL-0151, 0165, 0175, 0186.",
    },
    {
        "id": "Q8",
        "user": "michael",
        "new_chat": True,
        "question": "Show a heatmap of 2026 spills by geographic area and month.",
        "expected_note": "Expected area×month heatmap with corrected detections removed.",
    },
    {
        "id": "Q9",
        "user": "michael",
        "new_chat": True,
        "question": "Which rules has WELL-0001 broken in the last 12 months, and what is still open?",
        "expected_note": "Expected: SR-01/02/03, RC-04 breaches; USD 5,000 fines; pending RCA/records.",
    },
    {
        "id": "Q10",
        "user": "michael",
        "new_chat": True,
        "question": (
            "In Southwest Block 24, which spills in the last 90 days were not contained "
            "within the SOP time limit?"
        ),
        "expected_note": "Expected: WELL-0082 (~7.1 h), WELL-0100 (~8.1 h).",
    },
    # Re-seed Q6 privately before Q11 (Day-2 memory scenario); not a separate report section
    {
        "id": "Q6_seed",
        "user": "michael",
        "new_chat": True,
        "skip_doc": True,
        "question": (
            "Which 3 wells in Northeast Block 7 had the most downtime in the last 4 months? "
            "Show a bar chart."
        ),
        "expected_note": "Memory seed for Q11 (same as Q6).",
    },
    {
        "id": "Q11",
        "user": "michael",
        "new_chat": True,
        "question": (
            "For the three wells we discussed yesterday, what have repairs cost this year?"
        ),
        "expected_note": (
            "Expected (after Q6): WELL-0121 USD 25,794 · WELL-0138 USD 46,942 · WELL-0107 USD 5,176."
        ),
    },
    {
        "id": "Q12a",
        "user": "michael",
        "new_chat": True,
        "question": (
            "The 15 Oct reading for WELL-0164 is wrong. The handheld thermal gun used in that "
            "inspection was faulty; the field crew re-measured the wellhead at 41.0 °C."
        ),
        "expected_note": (
            "Expected: correction on INS-0164-0205; producing; disclose Pump jack conflict."
        ),
    },
    {
        "id": "Q12b",
        "user": "michael",
        "new_chat": True,
        "question": (
            "Which wells in Northeast Block 8 are not producing oil right now, and why? "
            "Show them in a table."
        ),
        "expected_note": "Expected after correction: WELL-0170, 0172, 0175 (not 0164).",
    },
    {
        "id": "Q13",
        "user": "michael",
        "new_chat": True,
        "question": (
            "The AI estimated the 24 Sep spill at WELL-0151 as 4.5 bbl. "
            "The field crew measured 5.4 bbl."
        ),
        "expected_note": "Expected: reportable SR-01 overdue; SR-02/SR-03 obligations updated.",
    },
    {
        "id": "Q14a",
        "user": "michael",
        "new_chat": True,
        "question": "Which wells in Northeast Block 7 need a root-cause analysis on their pump jack?",
        "expected_note": "Expected baseline: WELL-0108, 0121, 0138.",
    },
    {
        "id": "Q14b",
        "user": "michael",
        "new_chat": True,
        "question": (
            "Note for all T350 wells: our team now wants a pump jack root-cause analysis "
            "after 2 repairs, not 3."
        ),
        "expected_note": "Expected: save team rule for T350 threshold 2.",
    },
    {
        "id": "Q14c",
        "user": "michael",
        "new_chat": True,
        "question": "Which wells in Northeast Block 7 need a root-cause analysis on their pump jack?",
        "expected_note": "Expected with rule: WELL-0108, 0120, 0121, 0138 (not T200 WELL-0102).",
    },
    {
        "id": "Q15a",
        "user": "jessica",
        "new_chat": True,
        "question": "Is WELL-0164 producing oil right now? Show me its status card.",
        "expected_note": "Expected: yes via Michael's correction; mention Pump jack conflict.",
    },
    {
        "id": "Q15b",
        "user": "laura",
        "new_chat": True,
        "question": (
            "Which spills in Northeast Block 8 in the last 90 days needed a report to the regulator?"
        ),
        "expected_note": "Expected: WELL-0151 at 5.4 bbl overdue; no other ≥5 bbl.",
    },
    {
        "id": "Q16a",
        "user": "jake",
        "new_chat": True,
        "question": (
            "FYI for everyone: WELL-0007 will be shut in from 17 Oct for a 10-day planned workover."
        ),
        "expected_note": "Expected: shared note saved for everyone.",
    },
    {
        "id": "Q16b",
        "user": "michael",
        "new_chat": True,
        "question": "Is any downtime planned for WELL-0007?",
        "expected_note": "Expected: yes, according to Jake Morales.",
    },
    {
        "id": "Q16c",
        "user": "jessica",
        "new_chat": True,
        "question": (
            "For the three wells we discussed yesterday, what have repairs cost this year?"
        ),
        "expected_note": "Expected: no private history for Jessica (must not use Michael's).",
    },
    {
        "id": "Q17a",
        "user": "michael",
        "new_chat": True,
        "question": "WELL-0175 has been shut in for planned work since 14 Oct.",
        "expected_note": "Expected: record planned-work note aligned with inspection.",
    },
    {
        "id": "Q17b",
        "user": "tyler",
        "new_chat": True,
        "question": "WELL-0175 is running fine.",
        "expected_note": "Expected: unconfirmed; inspection shows not producing; matches Michael.",
    },
    {
        "id": "Q17c",
        "user": "jake",
        "new_chat": True,
        "question": "WELL-0001 is back online since 14 Oct.",
        "expected_note": "Expected: unconfirmed vs 15 Oct inspection / CTX-0001-003.",
    },
    {
        "id": "Q18",
        "user": "laura",
        "new_chat": True,
        "question": "Was the 20 Aug spill at WELL-0101 real oil?",
        "expected_note": (
            "Expected: yes, detection stands; show CTX disagreement; repair done; not formal corrections."
        ),
    },
]


def select_user(page, user_key: str) -> None:
    uid, name = USERS[user_key]
    page.evaluate(
        """([uid]) => {
          localStorage.setItem('fieldsense_user', uid);
          localStorage.setItem('fieldsense_as_of', '2026-10-17T00:00:00Z');
        }""",
        [uid],
    )
    page.reload(wait_until="networkidle")
    page.wait_for_timeout(800)
    sel = page.locator("select").first
    # Prefer label match
    options = sel.locator("option").all()
    for opt in options:
        text = opt.inner_text()
        if name in text:
            sel.select_option(value=opt.get_attribute("value") or uid)
            break
    else:
        sel.select_option(value=uid)
    page.wait_for_timeout(1000)


def new_chat(page) -> None:
    btn = page.get_by_role("button", name=re.compile(r"New chat", re.I))
    if btn.count():
        btn.first.click()
        page.wait_for_timeout(600)


def ask(page, question: str) -> None:
    box = page.locator("textarea").first
    box.fill(question)
    page.get_by_role("button", name=re.compile(r"^Send$", re.I)).click()
    # Wait for assistant response (loading disappears / answer panel appears)
    page.wait_for_timeout(500)
    page.wait_for_function(
        """() => {
          const t = document.body.innerText;
          return !t.includes('Calculating from inspections') &&
                 (t.includes('View evidence') || t.includes('WELL-') ||
                  t.includes('Shared') || t.includes('Correction') ||
                  t.includes('unconfirmed') || t.includes('days') ||
                  t.includes('readings') || t.includes('breach') ||
                  t.includes('I don') || t.includes("don't have"));
        }""",
        timeout=120_000,
    )
    page.wait_for_timeout(800)


def screenshot_answer(page, path: Path) -> None:
    # Prefer the last assistant bubble / main chat column
    main = page.locator("main").first
    main.screenshot(path=str(path))


def latest_answer_for_user(user_id: str) -> dict:
    """Read the newest assistant answer from the user's conversations (UI-driven runs)."""
    with httpx.Client(base_url=API, timeout=60.0) as client:
        client.post(
            "/api/v1/session/user",
            json={"user_id": user_id},
            headers={"X-User-Id": user_id},
        )
        convs = client.get("/api/v1/conversations", headers={"X-User-Id": user_id}).json()
        items = convs.get("conversations") or []
        if not items:
            return {}
        # API returns newest-first
        cid = items[0]["id"]
        detail = client.get(f"/api/v1/conversations/{cid}", headers={"X-User-Id": user_id}).json()
        for msg in reversed(detail.get("messages") or []):
            if msg.get("role") == "assistant" and msg.get("answer"):
                return msg["answer"]
            if msg.get("role") == "assistant" and msg.get("content"):
                return {"summary": msg["content"], "visuals": [], "rows": None}
    return {}


def build_docx(results: list[dict], out_path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

    title = doc.add_heading("FieldSense — Example Scenarios Q1–Q18", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(
        "Actual system answers captured from the running Docker UI "
        "(sample-v1, as-of 2026-10-17T00:00:00Z).\n"
        "Screenshots show the live FieldSense copilot response for each scenario step."
    )
    run.italic = True

    doc.add_heading("How to read this report", level=1)
    doc.add_paragraph(
        "Each section includes: (1) the scenario question as asked, (2) the official expected "
        "answer summary from the challenge brief, (3) the actual system summary returned by "
        "FieldSense, and (4) a browser screenshot of the UI answer (tables, charts, maps, cards)."
    )
    doc.add_paragraph(
        "Part A uses Michael Turner unless noted. Part B switches demo users. "
        "Q12–Q15 apply shared corrections; Q11/Q16 test private memory isolation."
    )

    for r in results:
        doc.add_page_break()
        doc.add_heading(f"{r['id']} — {r['user_name']}", level=1)

        doc.add_heading("Question asked", level=2)
        doc.add_paragraph(r["question"])

        doc.add_heading("Expected (challenge brief)", level=2)
        doc.add_paragraph(r["expected_note"])

        doc.add_heading("Actual system answer", level=2)
        summary = r.get("summary") or "(no summary)"
        doc.add_paragraph(summary)

        meta = doc.add_paragraph()
        meta.add_run("Visual / kind: ").bold = True
        kinds = []
        ans = r.get("answer") or {}
        for v in ans.get("visuals") or []:
            kinds.append(str(v.get("kind")))
        meta.add_run(", ".join(kinds) if kinds else ans.get("kind") or "short_answer / text")

        if ans.get("rows"):
            doc.add_heading("Result rows (system data)", level=2)
            rows = ans["rows"]
            if isinstance(rows, list) and rows and isinstance(rows[0], dict):
                keys = list(rows[0].keys())[:8]
                table = doc.add_table(rows=1 + min(len(rows), 20), cols=len(keys))
                table.style = "Table Grid"
                for i, k in enumerate(keys):
                    table.rows[0].cells[i].text = str(k)
                for ri, row in enumerate(rows[:20], start=1):
                    for ci, k in enumerate(keys):
                        val = row.get(k)
                        if isinstance(val, (dict, list)):
                            val = json.dumps(val)[:80]
                        table.rows[ri].cells[ci].text = "" if val is None else str(val)

        shot = Path(r["screenshot"]) if r.get("screenshot") else None
        if shot and shot.exists():
            doc.add_heading("UI screenshot", level=2)
            doc.add_picture(str(shot), width=Inches(6.5))
            cap = doc.add_paragraph(f"Figure: {r['id']} — FieldSense UI response")
            cap.runs[0].italic = True
            cap.runs[0].font.size = Pt(9)
            cap.runs[0].font.color.rgb = RGBColor(0x55, 0x65, 0x75)

    doc.add_page_break()
    doc.add_heading("Appendix — capture metadata", level=1)
    doc.add_paragraph(f"UI: {UI}")
    doc.add_paragraph(f"API: {API}")
    doc.add_paragraph("Analysis clock: 2026-10-17T00:00:00Z")
    doc.add_paragraph(f"Artifacts folder: {OUT}")
    doc.add_paragraph(f"Steps captured: {len(results)}")

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def main() -> None:
    SHOTS.mkdir(parents=True, exist_ok=True)
    results: list[dict] = []
    cookies = None

    with sync_playwright() as p:
        browser = p.chromium.launch(headless=True)
        context = browser.new_context(
            viewport={"width": 1440, "height": 1100},
            device_scale_factor=1.5,
        )
        page = context.new_page()
        page.goto(UI, wait_until="networkidle")
        page.wait_for_timeout(1000)

        current_user = None
        for sc in SCENARIOS:
            uid, uname = USERS[sc["user"]]
            print(f"=== {sc['id']} ({uname}) ===", flush=True)
            if current_user != sc["user"]:
                select_user(page, sc["user"])
                current_user = sc["user"]
            if sc.get("new_chat"):
                new_chat(page)

            ask(page, sc["question"])
            answer = latest_answer_for_user(uid)
            if sc.get("skip_doc"):
                (OUT / f"{sc['id']}.json").write_text(json.dumps(answer, indent=2, default=str))
                continue

            shot_path = SHOTS / f"{sc['id']}.png"
            screenshot_answer(page, shot_path)
            results.append(
                {
                    "id": sc["id"],
                    "user": sc["user"],
                    "user_name": uname,
                    "question": sc["question"],
                    "expected_note": sc["expected_note"],
                    "summary": answer.get("summary") or "",
                    "answer": answer,
                    "screenshot": str(shot_path),
                }
            )
            (OUT / f"{sc['id']}.json").write_text(json.dumps(answer, indent=2, default=str))

        browser.close()

    (OUT / "results.json").write_text(json.dumps(results, indent=2, default=str))
    docx_path = ROOT / "docs" / "FieldSense_Q1-Q18_System_Answers.docx"
    doc_alias = ROOT / "docs" / "FieldSense_Q1-Q18_System_Answers.doc"
    build_docx(results, docx_path)
    build_docx(results, doc_alias)
    print(f"Wrote {docx_path}")
    print(f"Wrote {doc_alias}")


if __name__ == "__main__":
    main()
